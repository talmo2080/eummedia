from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).parent
SOURCE = BASE / "웃자대한민국협회_홈페이지_디자인기획서.md"
TARGET = BASE / "웃자대한민국협회_홈페이지_디자인기획서.docx"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_run_font(run, name="Malgun Gothic", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_text(doc, text, style=None, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_run_font(run, size=10.5 if style is None else None, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    return p

def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, size={1: 17, 2: 13, 3: 11}[level], bold=True, color="000000")
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    p.paragraph_format.space_after = Pt(3)
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.7)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)
styles = doc.styles
styles["Normal"].font.name = "Malgun Gothic"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = title.add_run("웃자대한민국협회 홈페이지 디자인 기획서")
set_run_font(run, size=24, bold=True, color="000000")
subtitle = doc.add_paragraph()
subrun = subtitle.add_run("성창운 회장님 보고용 1차 기획안")
set_run_font(subrun, size=11, color="6B6B6B")
subtitle.paragraph_format.space_after = Pt(16)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line or line == ">" or line.startswith("# 웃자대한민국협회 홈페이지 디자인 기획서"):
        i += 1
        continue
    if line.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.2)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(line[2:])
        set_run_font(r, size=11, bold=True, color="806000")
        i += 1
        continue
    if line.startswith("## "):
        add_heading(doc, line[3:], 1); i += 1; continue
    if line.startswith("### "):
        add_heading(doc, line[4:], 2); i += 1; continue
    if line.startswith("#### "):
        add_heading(doc, line[5:], 3); i += 1; continue
    if line.startswith("| "):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[-: ]+", cell) for cell in row): rows.append(row)
            i += 1
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if ri == 0: set_cell_shading(cell, "F4E7BD")
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    r = p.add_run(value)
                    set_run_font(r, size=9.2, bold=(ri == 0))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue
    if line.startswith("- "):
        add_bullet(doc, line[2:]); i += 1; continue
    if re.match(r"^\d+\. ", line):
        add_number(doc, re.sub(r"^\d+\. ", "", line)); i += 1; continue
    if line.startswith("---"):
        i += 1; continue
    add_text(doc, line); i += 1

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("웃자대한민국협회 홈페이지 디자인 기획서")
set_run_font(fr, size=8, color="888888")
doc.save(TARGET)
print(TARGET)
