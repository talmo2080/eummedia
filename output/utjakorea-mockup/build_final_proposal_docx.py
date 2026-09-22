from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).parent
ASSET = BASE / "assets"
TARGET = BASE / "웃자대한민국협회_홈페이지_최종제안서_및_구축서.docx"

def font(run, size=10.5, bold=False, color="202020"):
    run.font.name = "Malgun Gothic"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Malgun Gothic")
    rpr.rFonts.set(qn("w:hAnsi"), "Malgun Gothic")
    rpr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

def para(doc, text="", size=10.5, bold=False, color="202020", align=None, space=5):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    r = p.add_run(text); font(r, size, bold, color)
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.line_spacing = 1.35
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text); font(r, {1:18, 2:14, 3:11.5}[level], True, "000000")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text); font(r, 10.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    return p

def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text); font(r, 10.5)
    p.paragraph_format.space_after = Pt(3)
    return p

def image(doc, path, width, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption); font(r, 8.5, False, "777777")
    cap.paragraph_format.space_after = Pt(10)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=""; shade(c,"F2E3B7"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r=c.paragraphs[0].add_run(h); font(r,9.3,True,"000000")
    for row in rows:
        cells=t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text=""; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            r=cells[i].paragraphs[0].add_run(value); font(r,9.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
doc.styles["Normal"].font.name="Malgun Gothic"; doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"),"Malgun Gothic"); doc.styles["Normal"].font.size=Pt(10.5)

title=doc.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.LEFT
r=title.add_run("웃자대한민국협회 홈페이지 최종 제안서 및 구축서"); font(r,25,True,"000000")
para(doc,"검색과 노출에 강하고 실시간 AI가 답하는 웃음 문화 플랫폼 구축안",12,False,"806000",space=2)
para(doc,"보고 대상: 성창운 회장님  |  작성 목적: 최종 방향 승인 및 협회 준비자료 확정",9,"000000"=="never", "777777", space=15)
image(doc, ASSET/"hero-smile-community.png", 6.4, "그림 1. 첫 화면 대표 이미지 시안: 여러 세대가 함께 웃는 장면과 문구 공간")

heading(doc,"1. 제안 요약",1)
para(doc,"웃자대한민국협회의 홈페이지는 단순한 행사 안내 페이지가 아니라, 검색을 통해 새로운 사람을 만나고 협회의 공식 정보를 AI가 정확히 안내하는 온라인 본부로 구축합니다.")
table(doc,["핵심 목표","구축 방향","기대 효과"],[
    ["검색과 노출","검색엔진이 읽기 쉬운 구조, 빠른 정적 페이지, 공식 정보의 일관된 관리","나비축제·힐링 프로그램·지역 활동을 검색으로 발견"],
    ["실시간 AI 안내","움직이는 스마일 클릭 → 실시간 AI 대화창 → 공식 자료 기반 답변","시간·장소·참가비·문의 방법을 즉시 안내"],
    ["브랜드 인상","실제 웃는 사람들의 사진, 따뜻한 색상, 짧은 문장, 큰 여백","행사 공지 모음이 아닌 살아 있는 문화협회로 인식"],
])

heading(doc,"2. 왜 새 홈페이지가 필요한가",1)
bullet(doc,"협회의 가장 큰 자산은 숫자보다 실제로 사람들이 만나고 웃는 현장입니다.")
bullet(doc,"검색 결과에서 협회의 공식 일정·장소·참가 방법이 한 번에 확인되어야 합니다.")
bullet(doc,"처음 방문한 사람도 30초 안에 ‘언제, 어디서, 어떻게 참여하는지’를 알아야 합니다.")
bullet(doc,"협회 정보가 바뀌었을 때 AI가 오래된 내용을 답하지 않도록 공식 자료를 한곳에서 관리해야 합니다.")
para(doc,"따라서 디자인과 기술을 따로 보지 않고, 사진·검색 구조·AI 안내·문의 동선을 하나의 시스템으로 설계합니다.")

heading(doc,"3. 첫 화면 디자인",1)
image(doc, BASE/"home-v2-desktop.png", 6.3, "그림 2. 현재 제작한 홈 화면 시안: 히어로, 핵심 정보 띠, 협회 소개, 프로그램 카드")
heading(doc,"3-1. 첫 화면의 핵심 문구",2)
para(doc,"누구나 와서, 함께 웃는 곳",22,True,"BD8200",space=3)
para(doc,"웃자대한민국협회는 매주 사람을 만나고, 웃음과 음악으로 일상을 밝힙니다. 예약 없이 편하게 오세요.")
heading(doc,"3-2. 화면 순서",2)
numbered(doc,"대표 히어로 사진 또는 짧은 무음 영상")
numbered(doc,"이번 주 핵심 정보: 매주 월요일 저녁 7시 / 624회 / 참가비 마음만!")
numbered(doc,"12년 활동과 협회 이야기")
numbered(doc,"이번 주 프로그램: 싱글벙글 나비축제·힐링 콘서트·힐링 테라피")
numbered(doc,"찾아가는 활동과 지역 이미지")
numbered(doc,"협회 사람들 및 성창운 회장님 소개")
numbered(doc,"오시는 길·전화·이메일 문의")

heading(doc,"4. 검색과 노출 특화 구축",1)
para(doc,"클로드가 제안한 검색·노출 중심 방향을 홈페이지의 기본 설계 원칙으로 반영합니다. 검색 노출은 특정 순위를 보장하는 일이 아니라, 검색엔진과 AI가 공식 정보를 정확히 읽고 연결할 수 있는 구조를 만드는 일입니다.")
table(doc,["영역","구축 내용","협회가 준비할 것"],[
    ["검색 기본 구조","페이지별 고유 제목·설명·주소, 의미 있는 H1/H2, 내부 링크","공식 페이지명과 핵심 소개 문구"],
    ["기사·행사 정보","행사명·일시·장소·문의·참가비를 실제 텍스트로 표시","최신 일정과 장소 확정본"],
    ["공유·미리보기","페이지별 OG 제목·설명·대표 이미지, 카카오 공유용 정보","페이지별 대표 이미지와 한 줄 설명"],
    ["검색 파일","sitemap.xml, robots.txt, canonical, 정적 HTML 또는 사전 렌더링","공개할 페이지 목록"],
    ["구조화 정보","협회·행사·장소 정보를 검색엔진이 이해할 수 있는 형태로 관리","공식 명칭·주소·연락처·행사 일정"],
    ["속도·접근성","이미지 압축, 지연 로딩, 대체 텍스트, 모바일 우선","사진 원본과 촬영·공개 동의"],
])

heading(doc,"4-1. 검색 노출을 위한 페이지 구성",2)
table(doc,["페이지","검색에 잡혀야 할 핵심","대표 행동"],[
    ["협회 이야기","웃자대한민국협회, 12년 활동, 시민 문화예술","협회 신뢰 확인"],
    ["싱글벙글 나비축제","매주 월요일 저녁 7시, 관악구, 참가비 마음만","이번 주 방문"],
    ["힐링 콘서트","산들소리, 음악, 야외무대, 일정","문의·참여"],
    ["힐링 테라피","웃음치료, 강사, 대상, 신청 방법","전화 문의"],
    ["찾아가는 활동","인생다방, 시장과 지역, 청년 웃음특강","기관 협업 문의"],
    ["오시는 길·문의","주소, 봉천역, 대표 전화, 이메일","길 찾기·전화"],
])

heading(doc,"5. 실시간 AI봇 구축",1)
para(doc,"최종 홈페이지에서는 스마일을 누르면 협회에 대해 무엇이든 물어볼 수 있고, 실시간 AI가 공식 자료를 바탕으로 답변하도록 구축합니다.")
heading(doc,"5-1. 사용자 흐름",2)
numbered(doc,"방문자가 오른쪽 아래에서 천천히 움직이는 스마일을 봅니다.")
numbered(doc,"스마일 또는 ‘뭐든 물어보세요!’ 말풍선을 클릭합니다.")
numbered(doc,"대화창에서 질문을 입력하거나 추천 질문을 누릅니다.")
numbered(doc,"브라우저가 보안된 서버 API에 질문을 전달합니다.")
numbered(doc,"AI가 협회의 공식 자료를 참고해 실시간 답변을 스트리밍합니다.")
numbered(doc,"근거가 없거나 확정되지 않은 질문에는 전화 문의 또는 공식 페이지를 안내합니다.")
heading(doc,"5-2. AI가 답할 내용",2)
bullet(doc,"이번 주 나비축제의 날짜·시간·장소")
bullet(doc,"힐링 콘서트와 힐링 테라피 일정·장소·참가 방법")
bullet(doc,"참가비와 예약 필요 여부")
bullet(doc,"봉천역에서 오는 길과 대표 전화")
bullet(doc,"찾아가는 활동과 기관 협업 문의 방법")
heading(doc,"5-3. AI 안전 원칙",2)
bullet(doc,"API 키는 브라우저에 넣지 않고 서버에서만 보관합니다.")
bullet(doc,"AI가 공식 자료에 없는 일정·가격·자격증·의료 정보를 임의로 단정하지 않도록 합니다.")
bullet(doc,"답변 기준 자료에 날짜와 최종 수정일을 기록합니다.")
bullet(doc,"잘못된 답변 신고 또는 전화 문의 연결을 제공합니다.")
bullet(doc,"질문 로그는 개인정보를 최소화하고 보관 기간을 정합니다.")

heading(doc,"6. 사이트 구조와 디자인",1)
table(doc,["메뉴","주요 화면","시각 방향"],[
    ["협회 이야기","12년 활동·사람들·협력기관·기록","골드·아이보리, 연혁과 인물 사진"],
    ["나비축제","매주 월요일 프로그램 안내","골드·크림, 참여자 현장 사진"],
    ["힐링 콘서트","산들소리와 음악·야외무대","블루·흰색, 숲과 무대 이미지"],
    ["힐링 테라피","웃음치료·강사·신청 안내","그린·크림, 강사와 자연"],
    ["찾아가는 활동","인생다방·시장·지역·청년 특강","레드·아이보리, 거리와 사람들"],
    ["오시는 길·문의","주소·지도·전화·이메일","골드·흰색, 가장 명확한 정보"],
])
image(doc, ASSET/"life-cafe.png", 3.0, "그림 3. 찾아가는 인생다방 활동 이미지")
image(doc, ASSET/"market-local.png", 4.7, "그림 4. 시장과 지역 활동 이미지")

heading(doc,"7. 협회가 준비해야 할 자료",1)
table(doc,["준비 자료","필요한 내용","용도"],[
    ["공식 기본정보","협회명·법인명·대표자·주소·전화·이메일","푸터·검색·AI 답변"],
    ["프로그램 정보","프로그램명·요일·시간·장소·참가비·예약 여부","상세 페이지·AI 답변"],
    ["사진","히어로·행사·인물·활동·장소 사진","홈·상세·공유 미리보기"],
    ["사진 권리","촬영일·촬영자·공개 동의·사용 범위","공개 안전성"],
    ["공식 답변","자주 묻는 질문과 확정 답변","AI 지식자료"],
    ["기록","행사 날짜·회차·대표 사진·설명","검색·연혁·갤러리"],
])
image(doc, ASSET/"seong-changwoon.jpg", 2.5, "그림 5. 협회장 소개 영역에 사용할 인물 이미지 예시")

heading(doc,"8. 제작 단계",1)
table(doc,["단계","작업","완료 기준"],[
    ["1단계 기획 승인","첫 화면·검색·AI 방향 확정","회장님과 협회 핵심 정보 승인"],
    ["2단계 디자인","홈·공통 헤더·푸터·모바일 시안","데스크톱·모바일 시각 검수"],
    ["3단계 콘텐츠","공식 사진·문구·프로그램 정보 입력","운영자 확인 완료"],
    ["4단계 검색 구축","메타·sitemap·OG·구조화 정보·성능","검색 기본 점검 통과"],
    ["5단계 AI 구축","서버 API·공식 지식자료·실시간 답변","질문 시 근거 있는 답변"],
    ["6단계 공개 전 점검","전화·지도·모바일·접근성·오답 점검","공개 승인"],
])

heading(doc,"9. 1차 공개 범위와 보류 범위",1)
heading(doc,"1차 공개 범위",2)
bullet(doc,"검색 가능한 홈과 6개 주요 안내 페이지")
bullet(doc,"대표 사진과 활동 기록")
bullet(doc,"전화·이메일·지도 연결")
bullet(doc,"스마일 AI의 실시간 질문·답변")
heading(doc,"보류 또는 후속 범위",2)
bullet(doc,"복잡한 회원가입·예약·결제 시스템")
bullet(doc,"검증되지 않은 방문자 수·검색 순위·도달 수치")
bullet(doc,"AI가 근거 없이 답할 수 있는 범용 질문 확대")
bullet(doc,"사진 공개 동의가 확인되지 않은 이미지의 공개")

heading(doc,"10. 승인 요청",1)
para(doc,"다음 네 가지를 승인하면, 승인된 범위 안에서 디자인과 구축을 이어갑니다.")
numbered(doc,"검색과 노출에 특화된 홈페이지로 구축한다.")
numbered(doc,"움직이는 스마일을 누르면 실시간 AI봇이 공식 협회 정보를 답한다.")
numbered(doc,"첫 화면은 실제 웃음 현장과 사람의 얼굴을 중심으로 만든다.")
numbered(doc,"협회는 공식 정보·사진·일정·AI 답변 자료를 준비한다.")

heading(doc,"11. 결론",1)
para(doc,"이 홈페이지는 협회의 활동을 예쁘게 보여주는 데서 끝나지 않습니다. 검색을 통해 새로운 참여자를 만나고, 스마일 AI를 통해 방문자의 질문에 바로 답하며, 공식 기록을 오래 남기는 온라인 협회 본부를 만드는 것이 목표입니다.")
para(doc,"클로드는 기획·콘텐츠·검색 구조를 검토하고, 코덱스는 디자인 구현·이미지 적용·브라우저 검수·기능 연결을 담당하는 방식으로 협업할 수 있습니다. 최종 기준은 이 제안서와 회장님이 승인한 공식 정보로 통일합니다.")

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fr=footer.add_run("웃자대한민국협회 홈페이지 최종 제안서 및 구축서"); font(fr,8,False,"888888")
doc.save(TARGET)
print(TARGET)
